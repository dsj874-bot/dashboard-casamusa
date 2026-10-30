"""Seguimiento diario de la venta Bticino + Legrand despues de la inversion
de stock (mercaderia recibida el 27 y 28-10-2026).

Objetivo (pedido del usuario, 2026-10-29): mostrar a Gerencia General que
con mas stock se vende mas AL CONTADO.

- "Antes" = dias habiles (lun-vie con venta en la empresa) del 01-09 al
  26-10-2026. "Despues" = desde el 29-10-2026 (primer dia completo con la
  mercaderia en tienda; 27 y 28 fueron dias de recepcion).
- Documentos de mas de $2 MM se tratan como proyecto: se listan aparte y
  NO entran a la comparacion (un solo proyecto, ej. Enersec $44,5 MM el
  21-10, distorsiona cualquier promedio).
- Control: el contado del resto de las marcas, para separar el efecto de
  la inversion de un dia/mes bueno en general.

Uso:  python scripts/seguimiento_inversion_lb.py [AAAA-MM-DD]
Sin fecha informa desde el ultimo dia habil con datos hasta el ultimo dia
cargado (el lunes incluye viernes y sabado). Ademas regenera el Excel con
la serie diaria completa.
"""
import sys, io, os, datetime as dt
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
import db

MARCAS = ["BTIC", "LEGR"]
INI_ANTES, FIN_ANTES = dt.date(2026, 9, 1), dt.date(2026, 10, 26)
INI_DESPUES = dt.date(2026, 10, 29)
UMBRAL_PROYECTO = 2_000_000
EXCEL = r"C:\Users\Marcelo\Desktop\Compras\Seguimiento Inversion Legrand\Seguimiento diario Bticino-Legrand.xlsx"
DIAS = ["lun", "mar", "mié", "jue", "vie", "sáb", "dom"]


def serie():
    """Una fila por dia con venta en la empresa desde INI_ANTES."""
    return db.query_all("""
      with d as (
        select fecha_conta f, doc_sap, marca = any(%(m)s) lb, coalesce(cond_pago,'') = 'CONTADO' cont,
               sum(total) v, sum(utilidad_bruta) u
        from v_ventas where fecha_conta >= %(ini)s group by 1,2,3,4)
      select f,
        coalesce(sum(v) filter (where lb and cont and abs(v) <= %(p)s),0) lb_cont,
        count(distinct doc_sap) filter (where lb and cont and v > 0 and abs(v) <= %(p)s) lb_docs,
        coalesce(sum(v) filter (where lb and not cont and abs(v) <= %(p)s),0) lb_cred,
        coalesce(sum(v) filter (where lb and abs(v) > %(p)s),0) lb_proy,
        coalesce(sum(v) filter (where lb),0) lb_tot,
        coalesce(sum(u) filter (where lb),0) lb_util,
        coalesce(sum(v) filter (where not lb and cont and abs(v) <= %(p)s),0) resto_cont
      from d group by f order by f""",
      {"m": MARCAS, "ini": INI_ANTES, "p": UMBRAL_PROYECTO})


def proyectos(desde, hasta):
    return db.query_all("""
      select fecha_conta f, doc_sap, nombre_cliente, sucursal_logica s, cond_pago, sum(total) v
      from v_ventas where fecha_conta between %(a)s and %(b)s and marca = any(%(m)s)
      group by 1,2,3,4,5 having abs(sum(total)) > %(p)s order by 1, 6 desc""",
      {"a": desde, "b": hasta, "m": MARCAS, "p": UMBRAL_PROYECTO})


def mm(x):
    return f"{float(x)/1e6:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def pct(a, b):
    if not b:
        return "—"
    v = (float(a) / float(b) - 1) * 100
    return f"{v:+.0f}%"


def promedio(filas, campo):
    return sum(float(r[campo]) for r in filas) / len(filas) if filas else 0.0


def main():
    filas = serie()
    habil = lambda r: r["f"].weekday() < 5
    antes = [r for r in filas if INI_ANTES <= r["f"] <= FIN_ANTES and habil(r)]
    despues = [r for r in filas if r["f"] >= INI_DESPUES and habil(r)]
    if not filas:
        print("Sin datos de ventas."); return

    ult = filas[-1]["f"]
    if len(sys.argv) > 1:
        ult = dt.date.fromisoformat(sys.argv[1])
    habiles_hasta = [r["f"] for r in filas if r["f"] <= ult and r["f"].weekday() < 5]
    desde = habiles_hasta[-1] if habiles_hasta else ult
    reporte = [r for r in filas if desde <= r["f"] <= ult]

    b = {k: promedio(antes, k) for k in ("lb_cont", "lb_docs", "lb_cred", "resto_cont")}
    print(f"BASE 'ANTES' ({len(antes)} días hábiles, {INI_ANTES:%d-%m} a {FIN_ANTES:%d-%m}): "
          f"contado {mm(b['lb_cont'])} MM/día en {b['lb_docs']:.0f} docs; crédito {mm(b['lb_cred'])} MM/día; "
          f"resto de marcas contado {mm(b['resto_cont'])} MM/día")
    print()
    print("DÍA(S) A INFORMAR")
    for r in reporte:
        mg = float(r["lb_util"]) / float(r["lb_tot"]) * 100 if r["lb_tot"] else 0
        sab = r["f"].weekday() >= 5
        print(f"- {DIAS[r['f'].weekday()]} {r['f']:%d-%m}: contado {mm(r['lb_cont'])} MM "
              f"({r['lb_docs']} docs){'' if sab else ' vs base ' + pct(r['lb_cont'], b['lb_cont'])}; "
              f"crédito {mm(r['lb_cred'])} MM{'' if sab else ' vs base ' + pct(r['lb_cred'], b['lb_cred'])}; "
              f"proyectos {mm(r['lb_proy'])} MM; total {mm(r['lb_tot'])} MM; margen {mg:.1f}%; "
              f"control resto contado {mm(r['resto_cont'])} MM"
              f"{'' if sab else ' (' + pct(r['resto_cont'], b['resto_cont']) + ')'}"
              f"{'  [sábado: sin base comparable]' if sab else ''}")
    for p in proyectos(desde, ult):
        print(f"  proyecto aparte: {p['f']:%d-%m} doc {p['doc_sap']} {p['nombre_cliente']} ({p['s']}, {p['cond_pago']}) {mm(p['v'])} MM")
    print()
    if despues:
        d = {k: promedio(despues, k) for k in ("lb_cont", "lb_docs", "lb_cred", "resto_cont")}
        print(f"ACUMULADO 'DESPUÉS' ({len(despues)} días hábiles desde {INI_DESPUES:%d-%m}): "
              f"contado {mm(d['lb_cont'])} MM/día ({pct(d['lb_cont'], b['lb_cont'])} vs antes), "
              f"{d['lb_docs']:.0f} docs/día ({pct(d['lb_docs'], b['lb_docs'])}); "
              f"crédito {mm(d['lb_cred'])} MM/día ({pct(d['lb_cred'], b['lb_cred'])}); "
              f"control resto contado {pct(d['resto_cont'], b['resto_cont'])}")
        share = sum(float(r["lb_cont"]) for r in despues) / max(1, sum(float(r["lb_cont"]) + float(r["lb_cred"]) for r in despues)) * 100
        share_a = sum(float(r["lb_cont"]) for r in antes) / max(1, sum(float(r["lb_cont"]) + float(r["lb_cred"]) for r in antes)) * 100
        print(f"PESO DEL CONTADO (sin proyectos): antes {share_a:.0f}% -> después {share:.0f}%")
    else:
        print(f"ACUMULADO 'DESPUÉS': aún sin días hábiles desde {INI_DESPUES:%d-%m}.")

    escribir_excel(filas, b)


def escribir_excel(filas, b):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Diario"
    cab = ["Fecha", "Día", "Periodo", "Contado B+L", "Docs contado", "Crédito B+L", "Proyectos (>2 MM)",
           "Total B+L", "Margen %", "Contado resto marcas", "Contado vs base", "Control resto vs base"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="22A347")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in filas:
        per = "Antes" if r["f"] <= FIN_ANTES else ("Recepción" if r["f"] < INI_DESPUES else "Después")
        hab = r["f"].weekday() < 5
        ws.append([r["f"], DIAS[r["f"].weekday()], per, float(r["lb_cont"]), r["lb_docs"], float(r["lb_cred"]),
                   float(r["lb_proy"]), float(r["lb_tot"]),
                   float(r["lb_util"]) / float(r["lb_tot"]) if r["lb_tot"] else None, float(r["resto_cont"]),
                   float(r["lb_cont"]) / b["lb_cont"] - 1 if hab and b["lb_cont"] else None,
                   float(r["resto_cont"]) / b["resto_cont"] - 1 if hab and b["resto_cont"] else None])
        if per == "Después":
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor="E3F4E8")
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "dd-mm-yyyy"
        for i in (3, 5, 6, 7, 9):
            row[i].number_format = '#,##0'
        for i in (8, 10, 11):
            row[i].number_format = '0.0%'
    for col, w in zip("ABCDEFGHIJKL", (12, 6, 11, 14, 9, 14, 14, 14, 9, 16, 11, 13)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws2 = wb.create_sheet("Criterio")
    for t in [
        "Marcas: Bticino (BTIC) y Legrand (LEGR). Montos netos.",
        f"Antes = días hábiles {INI_ANTES:%d-%m-%Y} a {FIN_ANTES:%d-%m-%Y}. Después = desde {INI_DESPUES:%d-%m-%Y}. 27 y 28-10 = recepción de la mercadería.",
        f"Documentos sobre ${UMBRAL_PROYECTO:,.0f} = proyecto; van en su columna y no entran a contado/crédito.".replace(",", "."),
        "Contado = condición de pago CONTADO. Crédito = cualquier cuenta corriente.",
        "Control = contado del resto de las marcas: si sube Bticino-Legrand y el resto no, el efecto es de la inversión.",
        f"Base diaria antes: contado B+L {b['lb_cont']:,.0f}; crédito B+L {b['lb_cred']:,.0f}; contado resto {b['resto_cont']:,.0f}.".replace(",", "."),
    ]:
        ws2.append([t])
    ws2.column_dimensions["A"].width = 120
    os.makedirs(os.path.dirname(EXCEL), exist_ok=True)
    try:
        wb.save(EXCEL)
        print(f"\nExcel actualizado: {EXCEL}")
    except PermissionError:
        print(f"\nNo se pudo actualizar el Excel (¿está abierto?): {EXCEL}")


if __name__ == "__main__":
    main()
